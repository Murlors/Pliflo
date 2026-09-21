"""Synthetic, original rendering fixtures; writes only into a new output directory.

Development-only requirements: python-docx, python-pptx, openpyxl, Pillow.
Usage: python scripts/generate-render-fixtures.py <new-system-temp-directory>
"""
import json
import io
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as SlideInches


root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=False)
image = Image.new("RGB", (640, 360), "#eef1e8")
draw = ImageDraw.Draw(image)
draw.rectangle((20, 20, 620, 340), outline="#246644", width=5)
draw.line((20, 20, 620, 340), fill="#333333", width=3)
image.save(root / "image space.png")

doc = Document()
doc.add_heading("Pliflo 合成回归 / Synthetic regression", 0)
for family in ["Arial", "Microsoft YaHei", "SimSun", "MissingPlifloFont"]:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"{family}: English 中文 e\u0301 ffi 👩‍💻")
    run.font.name = family
    run.font.size = Pt(14)
doc.add_paragraph().add_run("Bold 粗体").bold = True
doc.add_paragraph().add_run("Italic 斜体").italic = True
table = doc.add_table(rows=5, cols=3)
table.style = "Table Grid"
for r, row in enumerate(table.rows):
    for c, cell in enumerate(row.cells):
        cell.text = f"单元格 {r + 1}/{c + 1}"
doc.add_picture(str(root / "image space.png"), width=Inches(4))
doc.add_page_break()
doc.add_paragraph("PAGE TWO / 第二页：内容不能丢失。")
section = doc.add_section(WD_SECTION.NEW_PAGE)
section.page_width, section.page_height = Inches(10), Inches(7)
doc.add_paragraph("LANDSCAPE SECTION / 不同页面尺寸")
doc.save(root / "中文 mixed pages.docx")

# Original tiny test font, obfuscated according to ECMA-376; no commercial font.
embedded = Document()
run = embedded.add_paragraph().add_run("A中 A中")
run.font.name = "Pliflo Fixture"
run.font.size = Pt(24)
run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Pliflo Fixture")
memory = io.BytesIO()
embedded.save(memory)
with zipfile.ZipFile(memory) as archive:
    entries = {name: archive.read(name) for name in archive.namelist()}
w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
relationships = "http://schemas.openxmlformats.org/package/2006/relationships"
table = ET.fromstring(entries["word/fontTable.xml"])
font = ET.SubElement(table, f"{{{w}}}font", {f"{{{w}}}name": "Pliflo Fixture"})
ET.SubElement(font, f"{{{w}}}embedRegular", {f"{{{r}}}id": "plifloFont", f"{{{w}}}fontKey": "{00112233-4455-6677-8899-AABBCCDDEEFF}"})
entries["word/fontTable.xml"] = ET.tostring(table, encoding="utf-8", xml_declaration=True)
rels = ET.Element(f"{{{relationships}}}Relationships")
ET.SubElement(rels, f"{{{relationships}}}Relationship", {"Id": "plifloFont", "Type": r + "/font", "Target": "fonts/fixture.odttf"})
entries["word/_rels/fontTable.xml.rels"] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
font_bytes = bytearray((Path(__file__).resolve().parent.parent / "crates/cairo-replay/tests/fixtures/font-a.ttf").read_bytes())
mask = bytes.fromhex("00112233445566778899AABBCCDDEEFF")[::-1]
for i in range(32):
    font_bytes[i] ^= mask[i % 16]
entries["word/fonts/fixture.odttf"] = font_bytes
types = ET.fromstring(entries["[Content_Types].xml"])
ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default", {"Extension": "odttf", "ContentType": "application/vnd.openxmlformats-officedocument.obfuscatedFont"})
entries["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
with zipfile.ZipFile(root / "embedded.docx", "w", zipfile.ZIP_DEFLATED) as archive:
    for name, content in entries.items():
        archive.writestr(name, content)

slides = Presentation()
for number in range(2):
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(SlideInches(0.5), SlideInches(0.5), SlideInches(8), SlideInches(1))
    box.text = f"Slide {number + 1} / 中文演示 e\u0301"
    slide.shapes.add_picture(str(root / "image space.png"), SlideInches(1), SlideInches(2), width=SlideInches(6))
slides.save(root / "slides.pptx")

book = Workbook()
sheet = book.active
sheet.title = "宽表"
for column in range(1, 25):
    sheet.column_dimensions[get_column_letter(column)].width = 18 if column < 24 else 25
border = Border(*( [Side(style="thin", color="444444")] * 4 ))
for row in range(1, 86):
    for column in range(1, 25):
        cell = sheet.cell(row, column, f"R{row} C{column}" if column < 24 else f"LAST_COLUMN_{row}")
        cell.font = Font(name="Arial", size=11, bold=row == 1)
        cell.border = border
        if row == 1:
            cell.fill = PatternFill("solid", fgColor="DDEEDD")
sheet.cell(2, 1, "中文宽表")
second = book.create_sheet("Second")
second.column_dimensions["A"].width = 45
second.cell(1, 1, "Second sheet / 第二张工作表")
book.save(root / "wide.xlsx")

(root / "mixed.md").write_text(
    "# 中文 Markdown\n\nEnglish 中文 e\u0301 👩‍💻\n\n**Bold** and *italic*\n\n"
    "| 名称 | Value |\n| --- | --- |\n| 中文 | 123 |\n\n![local](image%20space.png)\n",
    encoding="utf-8",
)
(root / "corrupt.docx").write_bytes(b"Synthetic corrupt Office document")
(root / "original.json").write_text(json.dumps({
    "version": 1, "width": 612, "height": 792,
    "size": {"widthPt": 612, "heightPt": 792}, "commands": [], "unsupported": [],
}), encoding="utf-8")
print(root)
