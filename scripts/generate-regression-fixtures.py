"""Original stress/regression inputs; writes only to a new directory.

Usage: python scripts/generate-regression-fixtures.py <new-temp-directory>
Development dependencies: python-docx, python-pptx, openpyxl, Pillow, reportlab, pypdf.
No private documents or commercial fonts are used.
"""
from pathlib import Path
import json
import sys
from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as SlideInches
from openpyxl import Workbook
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from pypdf import PdfReader, PdfWriter

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=False)
expected = {}

image = Image.new("RGB", (4096, 2048), "white")
draw = ImageDraw.Draw(image)
for x in range(0, 4096, 128):
    for y in range(0, 2048, 128):
        draw.rectangle((x, y, x+127, y+127), fill=(x % 256, y % 256, 120))
image.save(root / "large-image.png")
image.save(root / "large-image.jpg", quality=90)
image.save(root / "image space.png")
(root / "local-images.md").write_text(
    "# LOCAL-START\n\nBefore ![checker](image%20space.png) AFTER-IMAGE\n\n"
    "[![linked checker](large-image.png)](https://example.invalid)\n\nLOCAL-END\n\n"
    "![REMOTE-OMITTED](https://example.invalid/image.png)\n\n![PARENT-OMITTED](%2e%2e/image.png)\n",
    encoding="utf-8")
expected["local-images.md"] = {"markers": ["LOCAL-START", "Before", "AFTER-IMAGE", "LOCAL-END", "REMOTE-OMITTED", "PARENT-OMITTED"], "minImages": 2}

doc = Document()
for i in range(12):
    if i:
        doc.add_page_break()
    doc.add_heading(f"DOCX-P{i:02d} 中文多页", 1)
    for family in ["Arial", "Microsoft YaHei", "MissingRegressionFont"]:
        run = doc.add_paragraph(f"{family}: 中文 English e\u0301 👩‍💻").runs[0]
        run.font.name = family
        run.font.size = Pt(12)
        run.bold = i % 2 == 0
        run.italic = i % 2 == 1
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"CELL-{i}-{r}-{c}"
    doc.add_picture(str(root / "large-image.png"), width=Inches(4))
doc.save(root / "twelve-pages.docx")
expected["twelve-pages.docx"] = {"pages": 12, "minImages": 12,
    "markers": [f"DOCX-P{i:02d}" for i in range(12)] +
               [f"CELL-{i}-{r}-{c}" for i in range(12) for r in range(3) for c in range(3)]}

ppt = Presentation()
for i in range(8):
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_textbox(SlideInches(.5), SlideInches(.3), SlideInches(9), SlideInches(1)).text = f"PPTX-S{i:02d} 中文 English e\u0301"
    slide.shapes.add_picture(str(root / "large-image.png"), SlideInches(1), SlideInches(1.5), width=SlideInches(8))
ppt.save(root / "eight-slides.pptx")
expected["eight-slides.pptx"] = {"pages": 8, "minImages": 8, "markers": [f"PPTX-S{i:02d}" for i in range(8)]}
for name in ["large-image.png", "large-image.jpg", "image space.png"]:
    expected[name] = {"pages": 1, "minImages": 1}

book = Workbook()
sheet = book.active
for r in range(1, 101):
    for c in range(1, 33):
        sheet.cell(r, c, f"R{r:03d}C{c:02d}")
book.save(root / "wide-100-rows.xlsx")
expected["wide-100-rows.xlsx"] = {"markers": [f"R{r:03d}C{c:02d}" for r in range(1, 101) for c in range(1, 33)]}

parts = ["# Long blocks / 中文分页"]
markers = []
for kind, count in [("PAR", 900), ("QUOTE", 300), ("CODE", 180), ("CELL", 400)]:
    tokens = [f"{kind}{i:04d}" for i in range(count)]
    markers.extend(tokens)
    if kind == "PAR":
        parts.append(" ".join(tokens))
    elif kind == "QUOTE":
        parts.append("> " + " ".join(tokens))
    elif kind == "CODE":
        parts.append("```\n" + "\n".join(tokens) + "\n```")
    else:
        parts.append("| Long cell | End |\n| --- | --- |\n| " + " ".join(tokens) + " | TABLE-END |")
(root / "long-blocks.md").write_text("\n\n".join(parts), encoding="utf-8")
expected["long-blocks.md"] = {"markers": markers + ["TABLE-END"]}

pdf = canvas.Canvas(str(root / "mixed-sizes.pdf"), pagesize=(612, 792))
for i, size in enumerate([(612, 792), (842, 595), (420, 595)]):
    pdf.setPageSize(size)
    pdf.drawString(36, size[1]-50, f"PDF-P{i}")
    pdf.showPage()
pdf.save()
expected["mixed-sizes.pdf"] = {"pages": 3, "markers": [f"PDF-P{i}" for i in range(3)]}
writer = PdfWriter()
writer.append(PdfReader(root / "mixed-sizes.pdf"))
writer.encrypt("synthetic-password")
with (root / "encrypted.pdf").open("wb") as output:
    writer.write(output)
(root / "broken.pdf").write_bytes(b"%PDF-1.7\nsynthetic broken PDF")
(root / "expected.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")
print(root)
